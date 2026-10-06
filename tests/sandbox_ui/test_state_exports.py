import io
import json
import unittest
from copy import deepcopy
from decimal import Decimal
from openpyxl import load_workbook
from decision_core.storage import normalize_storage
from decision_core.generation import normalize_generation
from sandbox_compute.generation_runner import run_generation
from sandbox_ui.common import initial
from sandbox_ui.state import quick_payload,quick_differences,Submitted
from sandbox_ui.exports import business_sheets,business_bytes,technical,workbook
from sandbox_display.business_s4 import language_defects,storage_snapshot,price_snapshot


class StateExportsTests(unittest.TestCase):
    def test_quick_advanced_equal(self):
        for case,norm in (("storage",normalize_storage),("generation",normalize_generation)):
            p=initial(case)
            self.assertEqual(norm(p),norm(quick_payload(case,p)))

    def test_total_used_switch_requires_notice(self):
        p=initial("storage");p["budget"]={"mode":"total_minus_used","total_cycles":"1","used_cycles":"0.4"}
        self.assertTrue(any("额度" in s for s in quick_differences("storage",p)))
        q=quick_payload("storage",p)
        self.assertEqual(q["budget"]["remaining_cycles"],"0.6")
        self.assertEqual(p["budget"]["mode"],"total_minus_used")
        r=normalize_storage(p)
        self.assertTrue(any("1.00／0.40" in s["内容"] for s in storage_snapshot(r)))
        self.assertFalse(Submitted(r,None,"advanced",None).matches(normalize_storage(q)))

    def test_percent_precision_not_float_drift(self):
        p=initial("storage");p["current_soc_ratio"]="0.65123456789";p["terminal_soc_ratio"]=p["current_soc_ratio"]
        self.assertEqual(normalize_storage(quick_payload("storage",p)).current_soc_ratio,Decimal(p["current_soc_ratio"]))

    def test_existing_contract_preserved(self):
        p=initial("generation");p["existing_contract_mwh"]="4321";p["existing_contract_price"]="200"
        q=quick_payload("generation",p)
        self.assertEqual(q["existing_contract_mwh"],"4321")
        self.assertIn("原合同单价",quick_differences("generation",p))

    def test_bad_grid_not_reinterpreted(self):
        p=initial("storage");p["interval_minutes"]=30
        self.assertTrue(quick_differences("storage",p))
        # Correct half-hour rows are not interpolated to quick's hourly grid.
        p["prices"]=p["prices"][:1];p["prices"][0]["interval_end"]="2026-10-06T17:30:00+08:00";p["wait_until"]=None
        normalize_storage(p)
        with self.assertRaises(ValueError):quick_payload("storage",p)

    def generation_saved(self,unknown=False):
        p=initial("generation")
        if unknown:p["mechanism"]={"mode":"unknown","ratio":None,"cap_mode":"unknown","remaining_cap_mwh":None,"price_yuan_per_mwh":None}
        r=normalize_generation(p);return Submitted(r,run_generation(r),"advanced",None)

    def test_business_all_prices_versions(self):
        s=self.generation_saved();sheets=business_sheets(s)
        self.assertEqual(len(sheets["本次价格判断"]),3)
        self.assertTrue(any(r["项目"]=="计算核心版本" for r in sheets["边界与解释"]))
        text=[x for rows in sheets.values() for r in rows for k,v in r.items() for x in (k,v)]
        self.assertEqual(language_defects(text),[])

    def test_unknown_export_retains_evaluable_diagnostics(self):
        s=self.generation_saved(True)
        sheets=business_sheets(s)
        self.assertEqual(len(sheets["机制未确认的诊断"]),6)
        self.assertEqual(sheets["机制未确认的诊断"][0]["机制金额（元）"],"暂缺依据")
        self.assertNotEqual(sheets["机制未确认的诊断"][0]["市场金额（元）"],"暂缺依据")

    def test_excel_text_no_formula(self):
        book=load_workbook(io.BytesIO(workbook({"本次条件":[{"说明":"=HYPERLINK(\"x\")"}]})))
        self.assertEqual(book.active["A2"].data_type,"s")
        book=load_workbook(io.BytesIO(business_bytes(self.generation_saved(True))))
        self.assertIn("机制未确认的诊断",book.sheetnames)

    def test_technical_scope(self):
        s=self.generation_saved();v=technical(s)
        self.assertEqual(v["domain"],"sandbox")
        self.assertTrue(v["amounts_recalculated"]);self.assertFalse(v["historical_result"])
        self.assertEqual(v["input"]["existing_contract_mwh"],"2000")
        with self.assertRaises(ValueError):technical(Submitted({},s.result,"quick",None))

    def test_changed_input_stales_result(self):
        s=self.generation_saved();p=s.request.editable_payload();p["new_contract_price"]="390"
        self.assertFalse(s.matches(normalize_generation(p)))
        self.assertFalse(s.matches(None))

    def test_no_mechanism_not_unknown(self):
        s=self.generation_saved();self.assertEqual(s.result.assessments[0].scenarios[0].evaluation.mechanism_difference.status.value,"not_applicable")
        self.assertEqual(price_snapshot(s.request)[0]["机制结算参考价（元/兆瓦时）"],"不适用")


if __name__=="__main__":unittest.main()

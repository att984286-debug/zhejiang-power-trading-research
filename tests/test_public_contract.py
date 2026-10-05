from copy import deepcopy
import gzip
import io
import json
from pathlib import Path
import shutil
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from public_ui.package import PublicPackage, ROOT, safe_file
from public_ui.schema import PublicDataError, FIELDS, RELEASE, canonical, validate
from public_ui.export import export_bytes, sheet
from public_ui.common import money, line_chart


@pytest.fixture(scope='module')
def package():
    return PublicPackage()


def test_explicit_schema_and_scalar_monitor():
    validate({'automatic_trade':False,'as_of':'2026-08-01T16:00:00'},'b_monitor')
    with pytest.raises(PublicDataError):
        validate({'observed_prefix_net_mwh':['1','2']},'b_monitor')
    with pytest.raises(PublicDataError):
        validate({'automatic_trade':False,'unknown_future_field':'unsafe'},'b_monitor')
    with pytest.raises(PublicDataError):
        validate({'as_of':['nested_not_allowed']},'b_monitor')


@pytest.mark.parametrize('value',[
    '/Users/'+'probe'+'/sensitive', '/home/'+'probe'+'/secret',
    'file:'+'//private/hidden', 'ghp_'+'A'*30, 'data/'+'private/'+'secret',
])
def test_security_probes_rejected(value):
    with pytest.raises(PublicDataError):validate({'source_id':value},'evidence')


def test_null_zero_negative_and_nonfinite():
    assert money(None)=='不可评价 / 未知'
    assert money('0')=='0.00'
    assert money('-15.01')=='-15.01'
    validate({'amount_yuan':None},'b_line')
    with pytest.raises(PublicDataError):validate({'amount_yuan':float('nan')},'b_line')


def test_all_registered_objects_and_coverage(package):
    assert package.all_valid()==205
    a=package.read('storage/index.json'); b=package.read('generation/index.json')
    assert len(a['rows'])==1600
    assert len({r['node_id'] for r in a['rows']})==4
    assert len(b['groups'])==3
    assert len(b['rules'])==36
    for name, spec in package.files.items():
        obj=package.read(name)
        if spec['kind']=='a_day':
            assert len(obj['strategies'])==8
            for strategy in obj['strategies']:
                assert len(strategy['actions'])==len(strategy['original_plan'])==96
                assert strategy['summary']['full_profit_yuan'] is None
                for row in strategy['actions']:
                    assert row['interval_minutes']==15
                    assert 0.1-1e-8<=float(row['soc_after_ratio'])<=0.9+1e-8
                assert strategy['summary']['reference_only']==('oracle' in strategy['strategy'])
        elif spec['kind']=='b_group':
            assert len(obj['strategies'])==4 and len(obj['stress'])==3
            if obj['private_details_omitted']:
                for strategy in obj['strategies']:
                    assert strategy['inputs']==strategy['submissions']==strategy['monitor']==[]
                    assert strategy['details_available'] is False
                    assert all(r['period']==strategy['decision']['month'] for r in strategy['lines'])
                assert all(case['lines']==[] for case in obj['stress'])
            else:
                for strategy in obj['strategies']:
                    assert len(strategy['inputs'])==31*48
                    assert len(strategy['submissions'])==31
                    assert all(len(s['points'])==96 for s in strategy['submissions'])
                    assert all(len(s['rows'])<=10 for c in obj['stress'] for s in c['lines'])


def test_private_branch_cannot_gain_inputs(package):
    group=deepcopy(package.read('generation/private_2026-09.json.gz'))
    group['strategies'][0]['inputs']=[{'q_net_mwh':'3'}]
    with pytest.raises(PublicDataError):validate(group,'b_group')
    group=deepcopy(package.read('generation/private_2026-09.json.gz'))
    group['strategies'][0]['lines'].append({'period':'2026-09-01T00:00','amount_yuan':'3'})
    with pytest.raises(PublicDataError):validate(group,'b_group')


def test_cross_view_native_sequences_absent(package):
    # All download/view data is this single projection; no alternate raw-data endpoint.
    for name, spec in package.files.items():
        obj=package.read(name)
        if spec['kind']=='a_day':
            for strategy in obj['strategies']:
                for row in strategy['actions']+strategy['original_plan']+strategy['explanations']:
                    assert not any(k.endswith('amount_yuan') or ('price' in k and k not in {'execution_price_target','future_actual_price_used'}) for k in row)
                    if 'future_actual_price_used' in row:assert isinstance(row['future_actual_price_used'],bool)
                assert all(not any('objective' in k or k.endswith('amount_yuan') for k in e) for e in strategy['events'])
        if spec['kind']=='b_group' and obj['private_details_omitted']:
            assert all(not s['inputs'] for s in obj['strategies'])


@pytest.mark.parametrize('bad',['../outside','/etc/passwd','a/../../outside','a\\outside'])
def test_path_traversal(bad,tmp_path):
    with pytest.raises(PublicDataError):safe_file(tmp_path,bad)


def test_symlink_escape(tmp_path):
    inner=tmp_path/'root';inner.mkdir();outside=tmp_path/'outside';outside.write_text('probe')
    (inner/'escape').symlink_to(outside)
    with pytest.raises(PublicDataError):safe_file(inner,'escape')


def copy_one_package(tmp_path):
    root=tmp_path/'release';(root/'config').mkdir(parents=True)
    for f in ('PUBLIC_PACKAGE_ANCHOR.json','public_profile.json'):
        shutil.copy2(ROOT/'config'/f,root/'config'/f)
    source=ROOT/'public_results'/RELEASE; target=root/'public_results'/RELEASE
    target.mkdir(parents=True);shutil.copy2(source/'manifest.json',target/'manifest.json')
    (target/'storage').mkdir();shutil.copy2(source/'storage/index.json',target/'storage/index.json')
    return root,target


def test_file_tampering_and_no_missing_fallback(tmp_path):
    root,target=copy_one_package(tmp_path);p=PublicPackage(root)
    (target/'storage/index.json').write_bytes(b'{}')
    with pytest.raises(PublicDataError):p.read('storage/index.json')
    with pytest.raises(PublicDataError):p.read('generation/index.json')
    with pytest.raises(PublicDataError):p.read('unregistered.json')


def test_manifest_tampering(tmp_path):
    root,target=copy_one_package(tmp_path)
    manifest=json.loads((target/'manifest.json').read_text());manifest['business_recalculated']=True
    (target/'manifest.json').write_text(canonical(manifest))
    with pytest.raises(PublicDataError):PublicPackage(root)


def test_formula_like_text_is_not_formula():
    book=Workbook();book.remove(book.active)
    sheet(book,'probe',[{'formula_or_reason':'=HYPERLINK("https://example.com","probe")','amount_yuan':'-15.01'}],'b_line')
    out=io.BytesIO();book.save(out);loaded=load_workbook(io.BytesIO(out.getvalue()),data_only=False)
    ws=loaded['probe'];assert ws['A2'].data_type=='s'
    assert ws['A2'].value.startswith('=HYPERLINK')
    assert ws['B2'].value==-15.01
    assert json.loads(ws['C2'].value)['amount_yuan']=='-15.01'


def check_export(case,strategy,payload,index):
    raw,metadata=export_bytes(case,strategy,payload,index)
    assert 0<len(raw)<8000000
    assert metadata['case']==case and metadata['raw_data_included'] is False
    assert strategy in metadata['selection_id']
    book=load_workbook(io.BytesIO(raw),data_only=False)
    for ws in book:
        for row in ws:
            for cell in row:assert cell.data_type!='f'
    exact=[]
    for ws in book:
        headers=[c.value for c in ws[1]]
        if 'exact_decimal_strings' in headers:
            i=headers.index('exact_decimal_strings')
            exact.extend(json.loads(row[i].value or '{}') for row in list(ws.rows)[1:])
    assert any(exact)
    if case=='B' and payload['private_details_omitted']:
        assert '合成申报接受' not in book.sheetnames and '合成日月明细' not in book.sheetnames
    # Workbook and JSON carry the original selected values; no regenerated business.
    if case=='A':
        risk=book['冻结风险']; assert risk.max_row>=2
        source_row=next(r for r in index['rows'] if r['node_id']==payload['node_id'] and r['market_date']==payload['market_date'] and r['strategy']==strategy)
        pack='august_2026' if payload['market_date'].startswith('2026-08') else 'counterfactual_2025'
        split='august_framework_only' if pack=='august_2026' else 'test'
        expected=[r['risk'] for r in index['comparisons'] if r['record']['group']==source_row['comparison_group'] and r['record']['strategy']==strategy and r['record']['split']==split and r['record']['pack']==pack]
        assert len(expected)==risk.max_row-1
        header=[c.value for c in risk[1]]; cvar=header.index('cvar_yuan')
        for i,row in enumerate(list(risk.rows)[1:]):
            if expected[i]['cvar_yuan'] is not None:assert abs(float(row[cvar].value)-float(expected[i]['cvar_yuan']))<1e-7
    return len(raw)


def test_all_strategy_export_selections(package):
    ai=package.read('storage/index.json');bi=package.read('generation/index.json')
    for node in sorted({r['node_id'] for r in ai['rows']}):
        for date in ['2025-09-02','2026-08-01']:
            day=package.read(f'storage/{node}/{date}.json.gz')
            for s in day['strategies']:check_export('A',s['strategy'],day,ai)
    for group in bi['groups']:
        payload=package.read(f"generation/{group['group']}.json.gz")
        for s in payload['strategies']:check_export('B',s['strategy'],payload,bi)


def test_plotly_no_payload_in_customdata_or_hover():
    fig=line_chart([('SOC',['2026-08-01'],['0.5'])],'SOC')
    spec=json.loads(fig.to_json())
    assert not spec['data'][0].get('customdata')
    assert not spec['data'][0].get('text')
    assert spec['data'][0]['y']==[0.5]

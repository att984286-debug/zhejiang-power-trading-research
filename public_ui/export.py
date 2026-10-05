"""Public exports use exactly the same filtered objects as the browser."""
from __future__ import annotations
import io
from decimal import Decimal,InvalidOperation
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from .schema import canonical,scalar,validate,PublicDataError
from .package import get_package

def export_metadata(case,selection,meta,evidence):
    validate(meta,'meta')
    for ref in evidence:validate(ref,'evidence')
    p=get_package()
    return {'schema_version':1,'case':case,'selection_id':selection,
            'release_id':p.anchor['release_id'],'public_manifest_sha256':p.anchor['manifest_sha256'],
            'publication_scope':meta['publication_scope'],'data_domain':meta['data_domain'],
            'information_status':meta['information_status'],'execution_assumption':meta['execution_assumption'],
            'component_coverage':meta['component_coverage'],'full_profit_status':meta['full_profit_status'],
            'amounts_recalculated':False,'raw_data_included':False,'source_evidence':evidence}

def sheet(book,name,rows,kind):
    for row in rows:validate(row,kind)
    ws=book.create_sheet(name)
    keys=list(dict.fromkeys(k for r in rows for k in r))
    if not keys:keys=['publication_status'];ws.append(keys);ws.append(['未发布此明细；不是0收益'])
    else:
        ws.append(keys+['exact_decimal_strings'])
        for row_index,row in enumerate(rows,2):
            exact={}
            for column,key in enumerate(keys,1):
                val=row.get(key)
                if isinstance(val,str) and any(t in key for t in ('yuan','mwh','_mw','ratio','coverage')):
                    try:
                        num=Decimal(val);exact[key]=val;val=num
                    except InvalidOperation:pass
                if isinstance(val,(list,dict)):val=canonical(val)
                if isinstance(val,str):scalar(val)
                cell=ws.cell(row_index,column,val)
                if isinstance(val,str):cell.data_type='s'
                if isinstance(val,(float,Decimal)):cell.number_format='#,##0.00;[Red](#,##0.00);0.00' if 'yuan' in key else '0.000'
            c=ws.cell(row_index,len(keys)+1,canonical(exact));c.data_type='s'
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='087F75');c.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
    for row in ws.iter_rows(min_row=2):
        for c in row:c.font=Font(name='Arial',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
    for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=28
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    return ws

def export_bytes(case,selection,payload,index):
    book=Workbook();book.remove(book.active)
    if case=='A':
        validate(payload,'a_day');validate(index,'storage_index')
        selected=next((r for r in payload['strategies'] if r['strategy']==selection),None)
        if selected is None:raise PublicDataError('导出选择不匹配')
        record=export_metadata('A',f"{payload['node_id']}|{payload['market_date']}|{selection}",payload['meta'],selected['evidence'])
        sheet(book,'场景与边界',[payload['meta']],'meta')
        sheet(book,'同日策略',[r['summary'] for r in payload['strategies']],'a_summary')
        sheet(book,'采用动作SOC',selected['actions'],'a_action')
        sheet(book,'原日前计划',selected['original_plan'],'a_action')
        sheet(book,'采用解释',selected['explanations'],'a_explain')
        sheet(book,'滚动事件',selected['events'],'a_event')
        sheet(book,'受控归因',[selected['attribution']],'a_attribution')
        source_row=next(r for r in index['rows'] if r['node_id']==payload['node_id'] and r['market_date']==payload['market_date'] and r['strategy']==selection)
        pack='august_2026' if payload['market_date'].startswith('2026-08') else 'counterfactual_2025'
        split='august_framework_only' if pack=='august_2026' else 'test'
        comparisons=[r for r in index['comparisons'] if r['record']['strategy']==selection and r['record']['group']==source_row['comparison_group']
                     and r['record']['pack']==pack and r['record']['split']==split]
        sheet(book,'冻结风险',[r['risk'] for r in comparisons],'a_risk')
        sheet(book,'证据索引',selected['evidence'],'evidence')
    elif case=='B':
        validate(payload,'b_group');validate(index,'generation_index')
        selected=next((r for r in payload['strategies'] if r['strategy']==selection),None)
        if selected is None:raise PublicDataError('导出选择不匹配')
        record=export_metadata('B',payload['group']+'|'+selection,payload['meta'],selected['evidence'])
        sheet(book,'场景与边界',[payload['meta']],'meta')
        sheet(book,'主体口径',[payload['identity']],'b_identity')
        sheet(book,'月前决策',[selected['decision']],'b_decision')
        sheet(book,'候选',[r['record'] for r in selected['candidates']],'b_candidate')
        sheet(book,'已评价分项',selected['coverage'],'b_coverage')
        sheet(book,'策略对照',payload['comparison'],'b_compare')
        if selected['details_available']:
            sheet(book,'合成申报接受',selected['inputs'],'b_input')
            sheet(book,'合成日月明细',selected['lines'],'b_line')
        else:sheet(book,'月科目范围',selected['lines'],'b_line')
        sheet(book,'风险',[selected['risk']],'b_risk')
        sheet(book,'分段差额',selected['chronological_rows'],'b_chron_row')
        sheet(book,'受控差额',[selected['attribution']],'b_attribution')
        sheet(book,'差额分项',selected['attribution_components'],'b_component')
        sheet(book,'预测与因子对照',selected['controls'],'b_control')
        sheet(book,'证据索引',selected['evidence'],'evidence')
    else:raise PublicDataError('未知导出Case')
    buf=io.BytesIO();book.save(buf)
    if buf.tell()>8000000:raise PublicDataError('公开导出体积超限')
    return buf.getvalue(),record

"""In-memory whitelisted business XLSX and independent technical JSON."""
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment

from decision_core.contracts import primitive
from decision_core.storage import StorageRequest
from decision_core.generation import GenerationRequest
from decision_core.storage_calculation import StorageCalculation
from decision_core.generation_calculation import GenerationCalculation
from decision_core.presets import preset_record
from decision_core.storage import normalize_storage
from decision_core.generation import normalize_generation
from sandbox_display.business_s4 import (storage_snapshot,generation_snapshot,action_rows,candidates,
    scenario_rows,amount,METHODS,safe_enum,price_snapshot,diagnostic_rows)

BOUNDARY="本次情景测算，不属于历史冻结研究结果；已评价收支不是完整利润。"


def workbook(sheets):
    book=Workbook();book.remove(book.active)
    for name,rows in sheets.items():
        ws=book.create_sheet(name)
        keys=list(dict.fromkeys(k for r in rows for k in r))
        if not keys:keys=["说明"];rows=[{"说明":"此范围无可展示记录；不是零金额"}]
        ws.append(keys)
        for row in rows:
            ws.append([str(row.get(key,"")) for key in keys])
        # All Chinese reader cells are text: no executable Excel formulas.
        for row in ws:
            for cell in row:
                cell.data_type="s"
                cell.font=Font(name="Arial",size=11)
                cell.alignment=Alignment(wrap_text=True,vertical="top")
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor="087F75")
            cell.font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=30
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    out=io.BytesIO();book.save(out);return out.getvalue()


def technical(saved):
    request,result=saved.request,saved.result
    if (type(request),type(result)) not in ((StorageRequest,StorageCalculation),(GenerationRequest,GenerationCalculation)):
        raise ValueError("导出对象不在沙盒白名单中")
    if result.output.metadata.input_sha256 != request.input_sha256 or result.historical_result or result.actual_execution_confirmed:
        raise ValueError("导出与本次输入或研究范围不一致")
    if saved.raw_input is not None and (normalize_storage if type(request) is StorageRequest else normalize_generation)(saved.raw_input) != request:
        raise ValueError("提交摘要与核心输入不一致")
    return {"schema_version":1,"domain":"sandbox","amounts_recalculated":True,
        "historical_result":False,"raw_research_data_included":False,
        "mode":saved.mode,"preset_id":saved.preset_id,
        "preset_record":preset_record("storage" if type(request) is StorageRequest else "generation") if saved.mode=="quick" else None,
        "submitted_input":saved.raw_input,"input":primitive(request),"result":primitive(result)}


def technical_bytes(saved):
    return json.dumps(technical(saved),ensure_ascii=False,allow_nan=False,indent=2).encode("utf-8")


def business_sheets(saved):
    technical(saved)
    request,result=saved.request,saved.result
    common={"边界与解释":[{"项目":"结果性质","内容":BOUNDARY},{"项目":"本次计算时间","内容":result.output.metadata.computed_at_utc.isoformat()},
        *[{"项目":"结果解释","内容":line} for line in result.advice.explanation_zh]],
        "本次条件":storage_snapshot(request) if type(request) is StorageRequest else generation_snapshot(request),
        "本次价格判断":price_snapshot(request)}
    common["边界与解释"] += [{"项目":"显示层版本","内容":"中文交互展示第一版"},
        {"项目":"计算核心版本","内容":"储能剩余策略第一版" if type(request) is StorageRequest else "发电月度头寸第一版"},
        {"项目":"输入方式","内容":"快速测算：标准研究预设" if saved.mode=="quick" else "高级情景测算：自主条件"},
        {"项目":"计算状态","内容":"有本次判断" if result.advice.available else "暂不能给出建议；查看诊断及技术记录"},
        {"项目":"金额精度","内容":"业务显示保留两位小数；技术记录保存原精度，逐行显示加总可能有舍入差异"}]
    if type(request) is StorageRequest:
        common["办法比较"]=[{"方案":METHODS[p.plan.method],"状态":safe_enum(p.plan.status.value),
            "预计收支差额（元）":amount(p.plan.contribution),"说明":p.diagnostic_zh} for p in result.computed_plans]
        for i,p in enumerate(result.computed_plans):
            common["动作明细"+str(i+1)]=action_rows(request,p)
    else:
        common["合同候选"]=candidates(result)
        if result.diagnostic_rows:common["机制未确认的诊断"]=diagnostic_rows(result)
        for i,a in enumerate(result.assessments):common["联合情况"+str(i+1)]=scenario_rows(a)
    return common


def business_bytes(saved):
    return workbook(business_sheets(saved))
